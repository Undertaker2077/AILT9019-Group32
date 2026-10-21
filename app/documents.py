import base64
import hashlib
import io
import os
from pathlib import Path, PurePosixPath

EXTENSIONS = {'.pdf', '.docx', '.txt', '.md'}
MAX_FILE = 2 * 1024 * 1024
MAX_FILES = 20
MAX_CHARS = 60000


def extract(name, raw):
    ext = Path(name).suffix.lower()
    if ext not in EXTENSIONS or raw is None or len(raw) > MAX_FILE:
        raise ValueError('Unsupported type or file exceeds 2 MB.')
    if ext == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(raw))
        if reader.is_encrypted or len(reader.pages) > 40:
            raise ValueError('Encrypted PDFs and PDFs over 40 pages are not supported.')
        units = [dict(locator=f'page {i+1}', text=page.extract_text() or '') for i, page in enumerate(reader.pages)]
    elif ext == '.docx':
        import zipfile
        from docx import Document
        with zipfile.ZipFile(io.BytesIO(raw)) as z:
            if sum(x.file_size for x in z.infolist()) > 20 * 1024 * 1024:
                raise ValueError('DOCX expanded size exceeds 20 MB.')
        doc = Document(io.BytesIO(raw))
        units = [dict(locator=f'paragraph {i+1}', text=p.text) for i, p in enumerate(doc.paragraphs) if p.text.strip()]
        for i, table in enumerate(doc.tables):
            for j, row in enumerate(table.rows):
                units.append(dict(locator=f'table {i+1}, row {j+1}', text=' | '.join(cell.text for cell in row.cells)))
    else:
        text = raw.decode('utf-8-sig')
        units = [dict(locator=f'line {i+1}', text=line) for i, line in enumerate(text.splitlines()) if line.strip()]
    if not any(x['text'].strip() for x in units):
        raise ValueError('No readable text. Scanned images need OCR, which is outside this version.')
    if sum(len(x['text']) for x in units) > MAX_CHARS:
        raise ValueError('Document text exceeds 60,000 characters; split the document.')
    return dict(filename=name, sha256=hashlib.sha256(raw).hexdigest(), units=units, size=len(raw))


def folder_files(folder):
    root = Path(folder).expanduser()
    if not root.is_dir() or root.is_symlink():
        raise ValueError('Select an existing local folder, not a symbolic link.')
    root = root.resolve()
    found = []
    for directory, dirs, files in os.walk(root, followlinks=False):
        dirs[:] = sorted(d for d in dirs if not d.startswith('.') and not (Path(directory)/d).is_symlink())
        for name in sorted(files):
            p = Path(directory)/name
            if name.startswith('.') or p.is_symlink() or p.suffix.lower() not in EXTENSIONS:
                continue
            try:
                p.resolve().relative_to(root)
            except ValueError:
                continue
            found.append(p)
            if len(found) > MAX_FILES:
                raise ValueError('More than 20 supported files. Select a smaller folder.')
    return [(p.relative_to(root).as_posix(), p.read_bytes() if p.stat().st_size <= MAX_FILE else None) for p in found]


def uploads(files):
    if not isinstance(files, list) or not 0 < len(files) <= MAX_FILES:
        raise ValueError('Select 1–20 files.')
    result = []
    for item in files:
        name = str(item['name']).replace('\\', '/')
        parts = PurePosixPath(name)
        if parts.is_absolute() or '..' in parts.parts or any(p.startswith('.') for p in parts.parts):
            raise ValueError('Invalid relative filename.')
        raw = base64.b64decode(item['content'], validate=True)
        if not raw or len(raw) > MAX_FILE:
            raise ValueError('Empty file or file exceeds 2 MB.')
        result.append((name, raw))
    return result
